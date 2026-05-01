"""
One-off script: regenerate DOCUMENTATION.docx from this file's content.
Run from project root: python scripts/build_documentation_docx.py
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt


def set_run_mono(run):
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(9)


def add_code_block(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(18)
    for line in text.strip().split("\n"):
        run = p.add_run(line + "\n")
        set_run_mono(run)


def main():
    root = Path(__file__).resolve().parent.parent
    out = root / "DOCUMENTATION.docx"

    doc = Document()
    title = doc.add_heading("Sunny Mart — Project documentation", 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    t = doc.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    t.rows[0].cells[0].text = "Student name"
    t.rows[0].cells[1].text = ""
    t.rows[1].cells[0].text = "Date"
    t.rows[1].cells[1].text = ""

    doc.add_paragraph()
    doc.add_heading("Abstract", level=1)
    doc.add_paragraph(
        "This project is a small e-commerce web application built for learning classic "
        "server-side web development. The storefront lets visitors browse products, manage "
        "a session-based shopping cart, and place orders that are stored in a SQLite database. "
        "A separate admin dashboard (protected by login) is used to create, edit, and delete "
        "products and to review orders. The user interface is split into plain HTML template "
        "files (with {{placeholder}} markers), CSS for styling, and a little JavaScript where "
        "helpful; PHP handles HTTP requests, database access, security-sensitive output escaping, "
        "and assembling the final HTML. The design keeps structure (templates), behavior (PHP), "
        "and presentation (CSS) relatively separate so it is easier to read and extend."
    )

    doc.add_heading("Table of contents", level=1)
    for line in [
        "1. How the project works",
        "2. How to run the project",
        "3. URLs and paths",
        "4. Folder and file structure",
        "5. Database",
        "6. Security notes (teaching points)",
    ]:
        doc.add_paragraph(line, style="List Number")

    doc.add_heading("1. How the project works", level=1)

    doc.add_heading("1.1 Overall flow", level=2)
    flow = [
        "The browser requests a PHP page (for example index.php).",
        "PHP loads includes/bootstrap.php, which starts the session, loads configuration, "
        "defines helpers (h, redirect, money_format_shop), and loads includes/template.php.",
        "Store pages also load includes/cart.php, which reads/writes the cart in $_SESSION['cart'].",
        "PHP talks to SQLite through config/database.php (Database::pdo()). The database file "
        "is created automatically on first use.",
        "PHP builds HTML by loading .html files from templates/ and replacing {{name}}-style "
        "placeholders via template(). Dynamic values should be passed through h() before "
        "insertion to reduce XSS risk.",
        "shop_header() / shop_footer() (or admin_header() / admin_footer()) print the shared "
        "layout; the middle of the page is filled by page-specific templates.",
        "The browser receives HTML and loads css/style.css and js/*.js (see section 3).",
    ]
    for item in flow:
        doc.add_paragraph(item, style="List Number")

    doc.add_heading("1.2 Storefront (public site)", level=2)
    t2 = doc.add_table(rows=4, cols=2)
    t2.style = "Table Grid"
    rows = [
        ("File", "Role"),
        ("index.php", "Lists active products; POST adds lines to the cart (respects stock)."),
        ("cart.php", "Shows cart; POST can update quantities or clear the cart."),
        (
            "checkout.php",
            "If the cart has items, shows a form; POST creates an order, order_items, "
            "decreases stock, clears the cart.",
        ),
    ]
    for i, (a, b) in enumerate(rows):
        t2.rows[i].cells[0].text = a
        t2.rows[i].cells[1].text = b

    doc.add_heading("1.3 Admin area", level=2)
    t3 = doc.add_table(rows=7, cols=2)
    t3.style = "Table Grid"
    admin_rows = [
        ("File", "Role"),
        ("admin/login.php", "Login form; sets session on success."),
        ("admin/logout.php", "Clears admin session."),
        ("admin/index.php", "Product list with edit/delete."),
        ("admin/product-form.php", "Create or edit a product."),
        ("admin/product-delete.php", "POST handler to delete a product."),
        ("admin/orders.php", "Lists orders; optional ?id= shows line items."),
    ]
    for i, (a, b) in enumerate(admin_rows):
        t3.rows[i].cells[0].text = a
        t3.rows[i].cells[1].text = b
    doc.add_paragraph(
        "includes/admin_auth.php defines admin_logged_in(), admin_require_login(), etc. "
        "Protected admin pages call admin_require_login() first."
    )

    doc.add_heading("1.4 Template placeholders", level=2)
    doc.add_paragraph(
        "Templates live under templates/shop/ and templates/admin/. Placeholders are replaced "
        "when calling template(), for example:"
    )
    add_code_block(
        doc,
        """<?php
template('shop/header.html', ['page_title' => h('Products'), ...]);
""",
    )
    doc.add_paragraph(
        "Some placeholders receive small HTML fragments built in PHP (e.g. table rows); "
        "those fragments must only contain data that PHP has already escaped where needed."
    )

    doc.add_heading("2. How to run the project", level=1)
    doc.add_heading("2.1 Requirements", level=2)
    doc.add_paragraph(
        "PHP 8+ with extensions: PDO, pdo_sqlite, session, password (default in most PHP builds)."
    )

    doc.add_heading("2.2 Option A — PHP built-in server", level=2)
    doc.add_paragraph("From the project root (folder that contains index.php):")
    add_code_block(doc, "php -S localhost:8080")
    doc.add_paragraph("Store: http://localhost:8080/")
    doc.add_paragraph("Admin: http://localhost:8080/admin/login.php")

    doc.add_heading("2.3 Option B — XAMPP (Apache)", level=2)
    for item in [
        "Copy the project folder into C:\\xampp\\htdocs\\ (e.g. htdocs\\gassaq\\).",
        "Start Apache in XAMPP Control Panel (MySQL is not required).",
        "Store: http://localhost/gassaq/ — Admin: http://localhost/gassaq/admin/login.php",
        "The data/ directory must be writable so SQLite can create/update shop.db.",
    ]:
        doc.add_paragraph(item, style="List Number")

    doc.add_heading("2.4 Default admin login", level=2)
    t4 = doc.add_table(rows=3, cols=2)
    t4.style = "Table Grid"
    for i, (a, b) in enumerate(
        [("Field", "Value"), ("Username", "admin"), ("Password", "admin123")]
    ):
        t4.rows[i].cells[0].text = a
        t4.rows[i].cells[1].text = b
    doc.add_paragraph(
        "If SHOP_ADMIN_PASSWORD is set in the environment before the database is first created, "
        "that value is used when seeding the admin user."
    )

    doc.add_heading("3. URLs and paths", level=1)
    doc.add_heading("3.1 Important paths on disk", level=2)
    t5 = doc.add_table(rows=8, cols=2)
    t5.style = "Table Grid"
    disk = [
        ("Path", "Meaning"),
        ("config/database.php", "Database bootstrap and schema creation."),
        ("includes/", "Shared PHP: bootstrap, templates, cart, admin auth."),
        ("templates/", "HTML layouts and fragments."),
        ("css/style.css", "Global styles."),
        ("js/", "Client scripts."),
        ("data/shop.db", "SQLite file (created at runtime)."),
        ("admin/", "Admin PHP entry points."),
    ]
    for i, (a, b) in enumerate(disk):
        t5.rows[i].cells[0].text = a
        t5.rows[i].cells[1].text = b

    doc.add_heading("3.2 URL paths in the browser", level=2)
    t6 = doc.add_table(rows=5, cols=2)
    t6.style = "Table Grid"
    urls = [
        ("URL (example)", "Script"),
        ("/index.php or /", "Store home"),
        ("/cart.php", "Cart"),
        ("/checkout.php", "Checkout"),
        ("/admin/login.php", "Admin login"),
    ]
    for i, (a, b) in enumerate(urls):
        t6.rows[i].cells[0].text = a
        t6.rows[i].cells[1].text = b

    doc.add_heading("3.3 Paths inside HTML templates", level=2)
    doc.add_paragraph(
        "Store templates use {{base_path}} = . so CSS is linked relative to the store root. "
        "Admin templates use {{root_path}} = .. so styles resolve to ../css/style.css."
    )
    doc.add_paragraph(
        "ROOT_PATH in bootstrap.php is the absolute filesystem path to the project root on the "
        "server (for loading template files), not the browser URL."
    )

    doc.add_heading("3.4 BASE_URL and redirects", level=2)
    doc.add_paragraph(
        "BASE_URL is derived from dirname($_SERVER['SCRIPT_NAME']) so redirect('cart.php') works "
        "whether the app is at the server root or in a subdirectory (e.g. /gassaq/)."
    )

    doc.add_heading("4. Folder and file structure", level=1)
    tree = r"""gassaq/
├── DOCUMENTATION.md
├── DOCUMENTATION.docx
├── index.php
├── cart.php
├── checkout.php
├── config/
│   └── database.php
├── includes/
│   ├── bootstrap.php
│   ├── template.php
│   ├── cart.php
│   └── admin_auth.php
├── templates/
│   ├── shop/
│   │   ├── header.html, footer.html, page-index.html
│   │   ├── product-card.html, product-buy-form.html
│   │   ├── page-cart.html, cart-empty.html, cart-filled.html, cart-row.html
│   │   └── page-checkout.html, checkout-form.html, checkout-success.html
│   └── admin/
│       ├── header.html, footer.html, login.html
│       ├── page-products.html, product-row.html, page-product-form.html
│       └── page-orders.html, order-row.html, order-detail.html, order-item-row.html
├── css/style.css
├── js/shop.js, admin.js
├── data/.gitignore, shop.db (runtime)
└── admin/*.php"""
    p = doc.add_paragraph()
    for line in tree.split("\n"):
        run = p.add_run(line + "\n")
        set_run_mono(run)

    doc.add_heading("5. Database", level=1)
    doc.add_paragraph("SQLite file: data/shop.db (data/ must exist and be writable).")
    t7 = doc.add_table(rows=5, cols=2)
    t7.style = "Table Grid"
    dbrows = [
        ("Table", "Purpose"),
        ("products", "Catalog: name, description, price, image URL, stock, active."),
        ("admins", "Admin username + password hash."),
        ("orders", "One row per checkout: customer, email, total, timestamp."),
        ("order_items", "Lines: order id, product id, qty, price at purchase."),
    ]
    for i, (a, b) in enumerate(dbrows):
        t7.rows[i].cells[0].text = a
        t7.rows[i].cells[1].text = b
    doc.add_paragraph(
        "Foreign keys: PRAGMA foreign_keys = ON. Old orders still list line items if a product "
        "was deleted (LEFT JOIN in admin)."
    )

    doc.add_heading("6. Security notes (teaching points)", level=1)
    for item in [
        "Output escaping: use h() for DB/form text in HTML unless intentionally outputting trusted HTML.",
        "Passwords: password_hash() / password_verify(), not plain text.",
        "Admin: admin_require_login(); use POST for destructive actions where applicable.",
        "Stock: checkout uses a transaction and checks stock on update.",
        "Production would add HTTPS, CSRF tokens, stricter validation, rate limiting, etc.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph()
    p_end = doc.add_paragraph()
    r_end = p_end.add_run("End of documentation.")
    r_end.italic = True

    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
